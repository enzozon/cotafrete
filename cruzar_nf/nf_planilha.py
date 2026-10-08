"""Grava a NF do portal numa coluna própria (N, "NF (MAESTRO)") da planilha dos estagiários.

A planilha é de gente: o robô não pode estragá-la nem passar por cima de quem
está trabalhando nela. Por isso:

1. planilha aberta no Excel (o Excel não deixa ninguém gravar enquanto está
   aberta) -> não grava e diz quem abriu (arquivo ~$ do Excel, se for recente);
2. lê o arquivo e guarda a assinatura (data + tamanho);
3. edição CIRÚRGICA: só as células da coluna N mudam, direto no XML da aba
   PEDIDOS; todas as outras partes do arquivo (estilos, outras abas, links,
   filtros) continuam byte a byte iguais. Não usa openpyxl para salvar: ele
   regrava a planilha inteira e perde formatação;
4. confere o arquivo novo ANTES de gravar: todas as outras células iguais e a
   coluna N com o esperado;
5. backup da planilha (os 30 últimos), grava num temporário na mesma pasta e,
   no último instante, confere de novo se está aberta ou se mudou; só então
   troca de uma vez (os.replace) — quem abrir depois vê a planilha nova inteira;
6. relê do disco e confere de novo.

Coluna N já usada para outra coisa -> não grava. Linha que casa com vários
pedidos do portal: vale o de mesmo número de PEDIDO; se ainda sobrar NF
diferente, não grava aquela linha (vai para "ambíguas").

    python -m cruzar_nf.nf_planilha [--gravar]

Compatível com Python 3.8.
"""

from __future__ import annotations

import argparse
import datetime as _dt
import glob
import io
import json
import os
import re
import shutil
import sys
import tempfile
import xml.etree.ElementTree as ET
import zipfile
from typing import Any, Callable, Dict, List, Optional, Tuple
from xml.sax.saxutils import escape

from cruzar_nf.cadastro_pedidos import COLUNA_NF_MAESTRO, _campo, _norm, ler_planilha_com_linhas
from cruzar_nf.planilha_portal import nf_de, parear

TITULO = COLUNA_NF_MAESTRO
ABA = "PEDIDOS"
COLUNA = "N"
COLUNA_ESTILO = "J"          # "Nº NOTA FISCAL": a coluna nova fica com a mesma cara
PASTA_BACKUPS = "backups_planilha"
MANTER_BACKUPS = 30
ARQ_TRAVA = "nf_planilha.trava"
NS_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
NS_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

ROW_RE = re.compile(r'<row\b[^>]*?\br="(\d+)"[^>]*?(?:/>|>.*?</row>)', re.S)
CELL_RE = re.compile(r'<c\b[^>]*?(?:/>|>.*?</c>)', re.S)
REF_RE = re.compile(r'\br="([A-Z]+)(\d+)"')
ESTILO_RE = re.compile(r'\bs="(\d+)"')


class GravacaoRecusada(RuntimeError):
    """Nada foi gravado na planilha; a mensagem diz por quê."""


# -- planilha aberta? -------------------------------------------------------------

def quem_abriu(caminho: str) -> Optional[str]:
    """Nome no arquivo ~$ que o Excel cria ao abrir. Só vale se o ~$ for mais novo que
    a planilha: um ~$ velho (Excel que caiu) diria o nome errado."""
    pasta, nome = os.path.split(caminho)
    for candidato in ("~$" + nome, "~$" + nome[2:]):
        arq = os.path.join(pasta, candidato)
        try:
            if os.path.getmtime(arq) < os.path.getmtime(caminho):
                continue
            with open(arq, "rb") as f:
                dados = f.read(200)
        except OSError:
            continue
        if dados:
            nome_usuario = dados[1:1 + dados[0]].decode("cp1252", "replace").strip()
            if nome_usuario:
                return nome_usuario
    return None


def esta_aberta(caminho: str) -> Optional[str]:
    """O Excel abre a planilha deixando os outros só lerem: se não dá para abrir para
    gravar, ela está aberta. O ~$ sozinho não basta (fica para trás quando o Excel cai)."""
    try:
        with open(caminho, "r+b"):
            return None
    except PermissionError:
        quem = quem_abriu(caminho)
        return ("A planilha está aberta no Excel" + (f" por {quem}" if quem else "") +
                ". Peça para salvar e fechar e tente de novo.")


# -- edição cirúrgica do xlsx -----------------------------------------------------

def _num_coluna(letras: str) -> int:
    n = 0
    for ch in letras:
        n = n * 26 + ord(ch) - 64
    return n


def parte_da_aba(dados: bytes, nome: str) -> str:
    """Caminho, dentro do xlsx, do XML da aba `nome` (ex.: xl/worksheets/sheet1.xml)."""
    with zipfile.ZipFile(io.BytesIO(dados)) as z:
        wb = ET.fromstring(z.read("xl/workbook.xml"))
        rid = next((s.get(f"{{{NS_REL}}}id") for s in wb.iter(f"{{{NS_MAIN}}}sheet")
                    if (s.get("name") or "").strip().upper() == nome), None)
        rels = ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
        alvo = next((r.get("Target") for r in rels if r.get("Id") == rid), None)
    if not alvo:
        raise GravacaoRecusada(f"A planilha não tem a aba {nome}.")
    return alvo.lstrip("/") if alvo.startswith("/") else "xl/" + alvo


def _celula(ref: str, estilo: Optional[str], valor: str) -> str:
    s = f' s="{estilo}"' if estilo else ""
    if re.fullmatch(r"[1-9]\d{0,14}", valor):       # número de verdade: sem triângulo verde no Excel
        return f'<c r="{ref}"{s}><v>{valor}</v></c>'
    return f'<c r="{ref}"{s} t="inlineStr"><is><t>{escape(valor)}</t></is></c>'


def editar_aba(xml: str, valores: Dict[int, str], coluna: str = COLUNA,
               coluna_estilo: str = COLUNA_ESTILO) -> str:
    """Põe `valores` (linha -> texto) na `coluna`, mexendo só nessas células."""
    alvo = _num_coluna(coluna)
    feitas = set()

    def trocar_linha(m: "re.Match[str]") -> str:
        numero = int(m.group(1))
        if numero not in valores:
            return m.group(0)
        linha = m.group(0)
        fim_tag = linha.index(">") + 1
        abertura = linha[:fim_tag]
        if abertura.endswith("/>"):
            abertura, corpo = abertura[:-2].rstrip() + ">", ""
        else:
            corpo = linha[fim_tag:-len("</row>")]
        celulas = CELL_RE.findall(corpo)
        if "".join(celulas) != corpo:
            raise GravacaoRecusada(f"Formato inesperado na linha {numero} da planilha; nada foi gravado.")
        colunas = [REF_RE.search(c).group(1) for c in celulas]
        estilos = {col: ESTILO_RE.search(c).group(1) for c, col in zip(celulas, colunas)
                   if ESTILO_RE.search(c)}
        estilo = estilos.get(coluna) or estilos.get(coluna_estilo)
        nova = _celula(f"{coluna}{numero}", estilo, valores[numero])
        if coluna in colunas:
            celulas[colunas.index(coluna)] = nova
        else:
            pos = next((i for i, col in enumerate(colunas) if _num_coluna(col) > alvo), len(celulas))
            celulas.insert(pos, nova)
        abertura = re.sub(r'\bspans="(\d+):(\d+)"',
                          lambda s: f'spans="{s.group(1)}:{max(int(s.group(2)), alvo)}"', abertura)
        feitas.add(numero)
        return abertura + "".join(celulas) + "</row>"

    novo = ROW_RE.sub(trocar_linha, xml)
    faltam = set(valores) - feitas
    if faltam:
        raise GravacaoRecusada(f"Não achei as linhas {sorted(faltam)[:5]} na planilha; nada foi gravado.")

    def dimensao(m: "re.Match[str]") -> str:
        inicio, col, lin = m.group(1), m.group(2), m.group(3)
        return f'<dimension ref="{inicio}{coluna if _num_coluna(col) < alvo else col}{lin}"/>'
    return re.sub(r'<dimension ref="([A-Z]+\d+:)([A-Z]+)(\d+)"/>', dimensao, novo, count=1)


def _reempacotar(original: bytes, parte: str, novo_xml: bytes) -> bytes:
    """Mesmo zip, mesma ordem, mesma compressão; só `parte` muda."""
    saida = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(original)) as zin, zipfile.ZipFile(saida, "w") as zout:
        for info in zin.infolist():
            zout.writestr(info, novo_xml if info.filename == parte else zin.read(info.filename))
    return saida.getvalue()


# -- conferência -----------------------------------------------------------------

def _texto(v: Any) -> str:
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return "" if v is None else str(v).strip()


def _titulo_aba(wb: Any) -> str:
    return next(t for t in wb.sheetnames if t.strip().upper() == ABA)


def conferir(antes: bytes, depois: bytes, esperado: Dict[int, str]) -> None:
    """Só a coluna N das linhas `esperado` pode ter mudado. Levanta GravacaoRecusada."""
    import openpyxl

    aba = parte_da_aba(antes, ABA)
    with zipfile.ZipFile(io.BytesIO(antes)) as za, zipfile.ZipFile(io.BytesIO(depois)) as zd:
        if [i.filename for i in za.infolist()] != [i.filename for i in zd.infolist()]:
            raise GravacaoRecusada("Na conferência, o arquivo novo não tem as mesmas partes; nada foi gravado.")
        for nome in za.namelist():
            if nome != aba and za.read(nome) != zd.read(nome):
                raise GravacaoRecusada(f"Na conferência, {nome} mudou sem querer; nada foi gravado.")
    alvo = _num_coluna(COLUNA) - 1
    wa = openpyxl.load_workbook(io.BytesIO(antes), read_only=True, data_only=True)
    wd = openpyxl.load_workbook(io.BytesIO(depois), read_only=True, data_only=True)
    try:
        la = list(wa[_titulo_aba(wa)].iter_rows(values_only=True))
        ld = list(wd[_titulo_aba(wd)].iter_rows(values_only=True))
    finally:
        wa.close()
        wd.close()
    if len(la) != len(ld):
        raise GravacaoRecusada("Na conferência, o número de linhas mudou; nada foi gravado.")
    for numero, (a, d) in enumerate(zip(la, ld), start=1):
        a, d = list(a), list(d)
        largura = max(len(a), len(d), alvo + 1)
        a += [None] * (largura - len(a))
        d += [None] * (largura - len(d))
        quer = esperado[numero] if numero in esperado else _texto(a[alvo])
        if a[:alvo] + a[alvo + 1:] != d[:alvo] + d[alvo + 1:] or _texto(d[alvo]) != quer:
            raise GravacaoRecusada(f"Na conferência, a linha {numero} não ficou como devia; nada foi gravado.")


# -- plano --------------------------------------------------------------------------

def planejar(linhas: List[Tuple[int, Dict[str, Any]]], pedidos: List[Dict[str, Any]]) -> Dict[str, Any]:
    pares = parear(linhas, pedidos)
    plano: Dict[str, Any] = {"mudancas": {}, "ambiguas": [], "sem_nf": 0, "iguais": 0, "atualizadas": 0,
                             "sem_par": 0}
    for n, l in linhas:
        if not pares[n]:
            plano["sem_par"] += 1
            continue
        candidatos = pares[n]
        mesmo_numero = [i for i in candidatos if _norm(_campo(l, "PEDIDO"))
                        and _norm(_campo(pedidos[i], "PEDIDO")) == _norm(_campo(l, "PEDIDO"))]
        no_portal = sorted({nf_de(pedidos[i]) for i in (mesmo_numero or candidatos)} - {""})
        if not no_portal:
            plano["sem_nf"] += 1
            continue
        if len(no_portal) > 1:
            plano["ambiguas"].append({"linha": n, "pedido": _texto(_campo(l, "PEDIDO")),
                                      "nfs": no_portal})
            continue
        atual = _texto(l.get(TITULO))
        if atual == no_portal[0]:
            plano["iguais"] += 1
            continue
        plano["mudancas"][n] = no_portal[0]
        plano["atualizadas"] += bool(atual)
    return plano


# -- gravação -----------------------------------------------------------------------

def _assinatura(caminho: str) -> Tuple[int, int]:
    st = os.stat(caminho)
    return st.st_mtime_ns, st.st_size


def _trocar(tmp: str, destino: str) -> None:
    os.replace(tmp, destino)


def _guardar_backup(dados: bytes, pasta: str) -> str:
    os.makedirs(pasta, exist_ok=True)
    arq = os.path.join(pasta, f"PLANILHA antes da NF {_dt.datetime.now():%Y-%m-%d_%H-%M-%S}.xlsx")
    with open(arq, "wb") as f:
        f.write(dados)
    for velho in sorted(glob.glob(os.path.join(pasta, "PLANILHA antes da NF *.xlsx")))[:-MANTER_BACKUPS]:
        os.remove(velho)
    return arq


def _ler(original: bytes) -> Tuple[List[Tuple[int, Dict[str, Any]]], Any]:
    """Linhas e o cabeçalho atual da coluna N, lidos de uma cópia local."""
    import openpyxl

    pasta = tempfile.mkdtemp(prefix="nf_planilha_")
    try:
        copia = os.path.join(pasta, "planilha.xlsx")
        with open(copia, "wb") as f:
            f.write(original)
        linhas = ler_planilha_com_linhas(copia)
        wb = openpyxl.load_workbook(copia, read_only=True)
        try:
            cab = next(wb[_titulo_aba(wb)].iter_rows(min_row=1, max_row=1, values_only=True), ())
        finally:
            wb.close()
        alvo = _num_coluna(COLUNA) - 1
        return linhas, (cab[alvo] if len(cab) > alvo else None)
    finally:
        shutil.rmtree(pasta, ignore_errors=True)


def gravar_nf_na_planilha(planilha: str, caminho_pedidos: str, pasta_dados: str, gravar: bool = False,
                          progresso: Callable[[str], None] = print,
                          _antes_de_trocar: Optional[Callable[[], None]] = None) -> Dict[str, Any]:
    from cruzar_nf.sincronizar import trava

    with trava(pasta_dados, "NF na planilha", arquivo=ARQ_TRAVA):
        return _rodar(planilha, caminho_pedidos, pasta_dados, gravar, progresso, _antes_de_trocar)


def _rodar(planilha: str, caminho_pedidos: str, pasta_dados: str, gravar: bool,
           progresso: Callable[[str], None], antes_de_trocar: Optional[Callable[[], None]]) -> Dict[str, Any]:
    from cruzar_nf.arquivo_pedidos import ArquivoPedidosConcorrente

    progresso("Conferindo se a planilha está aberta no Excel...")
    aberta = esta_aberta(planilha)
    if aberta and gravar:
        raise GravacaoRecusada(aberta)
    assinatura = _assinatura(planilha)
    with open(planilha, "rb") as f:
        original = f.read()
    if _assinatura(planilha) != assinatura:
        raise GravacaoRecusada("A planilha estava sendo salva enquanto eu lia; nada foi gravado. Tente de novo.")

    progresso("Lendo a planilha e os pedidos do portal...")
    linhas, cabecalho = _ler(original)
    if _texto(cabecalho) not in ("", TITULO):
        raise GravacaoRecusada(f"A coluna {COLUNA} da planilha já tem \"{cabecalho}\"; "
                               "não vou escrever por cima. Nada foi gravado.")
    manager = ArquivoPedidosConcorrente(caminho_pedidos)
    manager.iniciar()
    plano = planejar(linhas, manager.pedidos.get("PEDIDOS", []))
    mudancas = plano["mudancas"]
    res: Dict[str, Any] = {"gravou": False, "linhas": len(linhas), "a_gravar": len(mudancas),
                           "gravadas": 0, "atualizadas": plano["atualizadas"], "iguais": plano["iguais"],
                           "sem_nf": plano["sem_nf"], "sem_par": plano["sem_par"],
                           "ambiguas": plano["ambiguas"], "backup": None, "aviso": aberta,
                           "exemplos": [f"linha {n}: {nf}" for n, nf in list(mudancas.items())[:10]]}
    progresso(f"{len(mudancas)} NF(s) para gravar na coluna {COLUNA} ({plano['iguais']} já estão certas).")
    if not mudancas:
        return res

    valores = dict(mudancas)
    if _texto(cabecalho) != TITULO:
        valores[1] = TITULO
    aba = parte_da_aba(original, ABA)
    with zipfile.ZipFile(io.BytesIO(original)) as z:
        xml = z.read(aba).decode("utf-8")
    novo = _reempacotar(original, aba, editar_aba(xml, valores).encode("utf-8"))
    progresso("Conferindo a planilha nova antes de gravar...")
    conferir(original, novo, valores)
    if not gravar:
        return res

    res["backup"] = _guardar_backup(original, os.path.join(pasta_dados, PASTA_BACKUPS))
    tmp = os.path.join(os.path.dirname(os.path.abspath(planilha)), f"~nf_maestro_{os.getpid()}.tmp")
    try:
        with open(tmp, "wb") as f:
            f.write(novo)
        if antes_de_trocar:
            antes_de_trocar()
        aberta = esta_aberta(planilha)
        if aberta:
            raise GravacaoRecusada(aberta + " (abriram agora há pouco; nada foi gravado)")
        if _assinatura(planilha) != assinatura:
            raise GravacaoRecusada("A planilha mudou (alguém salvou) enquanto eu preparava; "
                                   "nada foi gravado. Tente de novo.")
        progresso("Gravando...")
        try:
            _trocar(tmp, planilha)
        except PermissionError:
            raise GravacaoRecusada("A planilha foi aberta no Excel agora há pouco; nada foi gravado. "
                                   "Tente de novo quando fecharem.")
    finally:
        if os.path.exists(tmp):
            os.remove(tmp)

    progresso("Conferindo a planilha gravada...")
    with open(planilha, "rb") as f:
        no_disco = f.read()
    if no_disco != novo:
        raise GravacaoRecusada("A conferência depois de gravar falhou: o arquivo no disco não é o que "
                               f"gravei (alguém salvou por cima?). Backup: {res['backup']}")
    conferir(original, no_disco, valores)
    res.update(gravou=True, gravadas=len(mudancas))
    progresso(f"Pronto: {len(mudancas)} NF(s) gravada(s) na coluna {COLUNA} ({TITULO}).")
    return res


def main(argv: Optional[List[str]] = None) -> int:
    from cruzar_nf.cadastro_pedidos import PLANILHA_PADRAO
    from cruzar_nf.sincronizar import DADOS_PADRAO, PEDIDOS_PADRAO

    try:
        from dotenv import load_dotenv
        load_dotenv()
    except ImportError:
        pass
    ap = argparse.ArgumentParser(prog="cruzar_nf.nf_planilha", description=__doc__.splitlines()[0])
    ap.add_argument("--planilha", default=os.environ.get("CADASTRO_PLANILHA") or PLANILHA_PADRAO)
    ap.add_argument("--dados", default=os.environ.get("SYNC_NF_DADOS") or DADOS_PADRAO)
    ap.add_argument("--pedidos", default=os.environ.get("SYNC_NF_PEDIDOS") or PEDIDOS_PADRAO)
    ap.add_argument("--gravar", action="store_true", help="grava na planilha (sem isso é só prévia)")
    a = ap.parse_args(argv)
    try:
        res = gravar_nf_na_planilha(a.planilha, a.pedidos, a.dados, gravar=a.gravar,
                                    progresso=lambda m: print(m, flush=True))
    except Exception as e:
        print(f"[ERRO] {e}", file=sys.stderr)
        return 1
    print(json.dumps(res, ensure_ascii=False, indent=1))
    return 0


if __name__ == "__main__":
    sys.exit(main())
