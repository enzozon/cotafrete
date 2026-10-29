"""Lê e valida a PLANILHA-BASE-v3 (aba ORÇAMENTO) sem acessar o HSE."""
import re
import unicodedata
import warnings
from dataclasses import dataclass
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

import openpyxl

LINHA_INI, LINHA_FIM = 2, 25
MAX_OBS = 1024
MAX_DESCRICAO = 512  # maxLength do #dsObsItem_<cdItem> no HSE (visto em 29/10/2026)
FRETES = ("CIF", "FOB")
FILIAIS = {1, 2, 3}
ICMS_PADRAO = Decimal("17")
CENTAVO = Decimal("0.01")

# A planilha usa listas suspensas x14; só lemos, então o aviso do openpyxl é ruído.
warnings.filterwarnings("ignore", message="Data Validation extension")

# Textos da aba LISTAS (e códigos do HSE) → código de rfFormaPg
FORMAS_PG = {
    "BOLETO": "BO", "CARTAO CREDITO": "CC", "CARTAO DE CREDITO": "CC",
    "CARTAO DEBITO": "CD", "CARTAO DE DEBITO": "CD", "CREDIARIO": "CA",
    "CREDIARIO (CARNE)": "CA", "DEPOSITO": "DE", "DINHEIRO": "DI", "PIX": "PX",
    **{c: c for c in ("BO", "CC", "CD", "CA", "DE", "DI", "PX")},
}


@dataclass(frozen=True)
class Item:
    numero: int            # Item 1 = linha 2
    linha: int
    produto: str
    quantidade: Decimal
    custo: Decimal
    valor_unit: Decimal    # coluna J, 2 casas
    codigo: str | None
    ncm: str | None
    origem: int
    icms: Decimal
    descricao_cadastro: str
    grupo: str | None
    marca: str | None
    descricao: str = ""  # coluna R: vai para a descrição do item no DAV


@dataclass(frozen=True)
class Orcamento:
    arquivo: Path
    itens: tuple[Item, ...]
    observacoes: str
    local_entrega: str | None
    cliente: str
    prazo_dias: int | None
    vendedor: str
    forma_pg: str | None
    parcelamento: str | None
    ordem_compra: str | None
    filial: int
    avisos: tuple[str, ...]
    frete: str | None = None  # B38: CIF ou FOB


class ErroPlanilha(Exception):
    def __init__(self, erros: list[str]):
        super().__init__("; ".join(erros))
        self.erros = erros


def normalizar(txt) -> str:
    if txt is None:
        return ""
    s = unicodedata.normalize("NFKD", str(txt))
    s = "".join(c for c in s if not unicodedata.combining(c))
    return " ".join(s.upper().split())


def _texto(v) -> str | None:
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip()
    return s or None


def _num(v) -> Decimal | None:
    """Número da célula; None se vazia; ValueError se não for número."""
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, bool):
        raise ValueError(v)
    if isinstance(v, (int, float)):
        return Decimal(str(v))
    s = str(v).strip().replace("R$", "").replace(" ", "")
    if "," in s or re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):  # "1.500" em texto = mil e quinhentos
        s = s.replace(".", "").replace(",", ".")
    try:
        return Decimal(s)
    except InvalidOperation:
        raise ValueError(v) from None


def _aba(wb):
    for ws in wb.worksheets:
        if normalizar(ws.title) == "ORCAMENTO":
            return ws
    raise ErroPlanilha(["aba ORÇAMENTO não encontrada"])


def linha_frete(frete: str, prazo: int) -> str:
    return f"FRETE {frete}, PRAZO DE ENTREGA {prazo} {'DIA' if prazo == 1 else 'DIAS'}"


def _com_frete(obs: str, frete: str | None, prazo: int | None) -> str:
    if not frete or not prazo:
        return obs
    return "\n".join(filter(None, (obs, linha_frete(frete, prazo))))


def observacoes_hse(orc: Orcamento) -> str:
    """Texto do #txObservacao: observações da B26 e, numa linha abaixo, o frete."""
    return _com_frete(orc.observacoes, orc.frete, orc.prazo_dias)


def _ler_item(ws, linha: int, erros: list[str], avisos: list[str]) -> Item | None:
    cel = {c: ws[f"{c}{linha}"].value for c in "ABCDEFGHIJKLMNOPQR"}
    produto, codigo = _texto(cel["A"]), _texto(cel["K"])
    if not produto and not codigo:
        return None
    n = linha - 1
    pre = f"Item {n}:"
    qtd_errada = len(erros)

    def num(col: str, nome: str) -> Decimal | None:
        try:
            return _num(cel[col])
        except ValueError:
            erros.append(f"{pre} {nome} inválido(a) ({cel[col]!r})")
            return None

    qtd = num("B", "quantidade")
    if len(erros) == qtd_errada and (qtd is None or qtd <= 0):
        erros.append(f"{pre} quantidade deve ser maior que zero")
    custo, d, e, f, g = (num(c, nome) for c, nome in
                         (("C", "CUSTO"), ("D", "FRETE PARA VIR"), ("E", "FRETE DELLA"),
                          ("F", "MADEIRA"), ("G", "MARGEM")))
    j = num("J", "VALOR DE VENDA UNITARIO")
    # a caixa do HSE tem uma linha só: quebra de linha (Alt+Enter) vira um espaço
    descricao = re.sub(r"\s*[\r\n]+\s*", " ", _texto(cel["R"]) or "")
    if len(descricao) > MAX_DESCRICAO:
        erros.append(f"{pre} DESCRIÇÃO NO DAV com {len(descricao)} caracteres (máximo {MAX_DESCRICAO})")
    if len(erros) > qtd_errada:
        return None

    custo = custo or Decimal(0)
    if j is None:  # arquivo salvo sem cache de fórmulas: mesma conta da planilha
        h = custo * qtd + (d or 0) + (e or 0) + (f or 0)
        i = h * (g or 0) / 100 + h
        j = i / qtd
        avisos.append(f"{pre} valor recalculado (planilha sem cache de fórmulas)")
    valor = j.quantize(CENTAVO, ROUND_HALF_UP)
    if valor <= 0:
        erros.append(f"{pre} valor unitário (coluna J) deve ser maior que zero")

    if codigo:  # com CÓDIGO HSE as colunas L–Q são ignoradas
        return Item(numero=n, linha=linha, produto=produto or "", quantidade=qtd, custo=custo,
                    valor_unit=valor, codigo=codigo, ncm=None, origem=0, icms=ICMS_PADRAO,
                    descricao_cadastro=produto or "", grupo=None, marca=None, descricao=descricao)

    ncm = _texto(cel["L"])
    if ncm:
        ncm = re.sub(r"\D", "", ncm)
        if len(ncm) != 8:
            erros.append(f"{pre} NCM deve ter 8 dígitos ({cel['L']!r})")
    origem = num("M", "ORIGEM")
    if origem is not None and (origem != origem.to_integral() or not 0 <= origem <= 8):
        erros.append(f"{pre} ORIGEM deve ser de 0 a 8 ({cel['M']!r})")
    icms = num("N", "ICMS")
    if icms is not None and 0 < icms < 1:  # célula formatada como %
        icms *= 100

    return Item(
        numero=n, linha=linha, produto=produto or "", quantidade=qtd, custo=custo,
        valor_unit=valor, codigo=codigo, ncm=ncm or None,
        origem=int(origem or 0), icms=icms if icms is not None else ICMS_PADRAO,
        descricao_cadastro=_texto(cel["O"]) or produto or "",
        grupo=_texto(cel["P"]), marca=_texto(cel["Q"]), descricao=descricao,
    )


def _filial(v, erros: list[str]) -> int:
    s = _texto(v)
    if not s:
        return 1
    m = re.match(r"\s*(\d+)", s)
    if not m or int(m.group(1)) not in FILIAIS:
        erros.append(f"Filial inválida ({s!r}); use 1, 2 ou 3")
        return 1
    return int(m.group(1))


def ler_planilha(path: Path) -> Orcamento:
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception as e:  # arquivo corrompido / não é xlsx
        raise ErroPlanilha([f"não foi possível abrir a planilha: {e}"]) from e
    ws = _aba(wb)
    erros: list[str] = []
    avisos: list[str] = []

    itens = [it for linha in range(LINHA_INI, LINHA_FIM + 1)
             if (it := _ler_item(ws, linha, erros, avisos))]
    if not itens and not erros:
        erros.append("nenhum item preenchido (linhas 2 a 25, coluna A ou K)")

    obs = _texto(ws["B26"].value) or ""
    cliente = _texto(ws["B31"].value)
    if not cliente:
        erros.append("Cliente (B31) é obrigatório")
    vendedor = _texto(ws["B33"].value)
    if not vendedor:
        erros.append("VENDEDOR (B33) é obrigatório")

    prazo = None
    try:
        p = _num(ws["B32"].value)
        if p is not None:
            if p < 0 or p != p.to_integral():
                raise ValueError(p)
            prazo = int(p)
    except ValueError:
        erros.append(f"Prazo de entrega (B32) inválido ({ws['B32'].value!r})")

    forma = _texto(ws["B34"].value)
    if forma:
        codigo = FORMAS_PG.get(normalizar(forma))
        if not codigo:
            erros.append(f"Forma de pagamento (B34) desconhecida ({forma!r})")
        forma = codigo

    frete = normalizar(ws["B38"].value) or None
    if frete and frete not in FRETES:
        erros.append(f"FRETE (B38) deve ser CIF ou FOB ({ws['B38'].value!r})")
    elif frete and not prazo:
        erros.append("Prazo de entrega (B32) é obrigatório quando o FRETE (B38) está marcado")
    texto = _com_frete(obs, frete if frete in FRETES else None, prazo)
    if len(texto) > MAX_OBS:
        com = " + linha do FRETE" if texto != obs else ""
        erros.append(f"OBSERVAÇÕES{com} com {len(texto)} caracteres (máximo {MAX_OBS})")

    filial = _filial(ws["B37"].value, erros)
    if erros:
        raise ErroPlanilha(erros)
    return Orcamento(
        arquivo=Path(path), itens=tuple(itens), observacoes=obs,
        local_entrega=_texto(ws["B30"].value), cliente=cliente, prazo_dias=prazo,
        vendedor=vendedor, forma_pg=forma, parcelamento=_texto(ws["B35"].value),
        ordem_compra=_texto(ws["B36"].value), filial=filial, avisos=tuple(avisos), frete=frete,
    )
