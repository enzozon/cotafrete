"""Sincronização bidirecional sobre arquivos sintéticos, nunca produção."""
import datetime as dt
import io
import json
import zipfile

import openpyxl
import pytest
from openpyxl.styles import Font
from openpyxl.worksheet.table import Table, TableStyleInfo

from cruzar_nf import sync_planilha_portal as SP
from cruzar_nf.cadastro_pedidos import cadastrar, chave
from cruzar_nf.arquivo_pedidos import ArquivoPedidosConcorrente
from cruzar_nf.nf_planilha import GravacaoRecusada, parte_da_aba

CAB = ['CIDADE', 'FRETE', 'PEDIDO', 'VALOR ', 'PRODUTO', 'DATA DE ENTREGA',
       'NMR DA RFQ', 'FATURAMENTO', 'STATUS', 'Nº NOTA FISCAL', 'DAV',
       'REQUISITANTE ', 'EMAIL REQUSITAN', 'NF (MAESTRO)']


def pedido(n=1, **kw):
    p = dict(zip(CAB, ['ITABIRA', 'EXW', n, 100, 'ITEM %s' % n, '10/11/2026',
                      n + 100, None, 'ABERTO', None, None, 'ANA', 'ana@vale.com', None]))
    p.update(kw)
    return p


class Cenario:
    def __init__(self, pasta, tabela=False):
        self.xlsx = pasta / 'planilha.xlsx'
        self.json = pasta / 'PEDIDOS.json'
        self.dados = pasta / 'dados'
        self.dados.mkdir()
        w = openpyxl.Workbook()
        s = w.active
        s.title = 'PEDIDOS'
        s.append(CAB)
        s.append(list(pedido().values()))
        for c in s[2]:
            c.font = Font(bold=True)
        s['F2'].number_format = 'dd/mm/yyyy'
        s.auto_filter.ref = 'A1:N2'
        if tabela:
            t = Table(displayName='Pedidos', ref='A1:N2')
            t.tableStyleInfo = TableStyleInfo(name='TableStyleMedium9')
            s.add_table(t)
        w.create_sheet('OUTRA')['A1'] = '=1+2'
        w.save(self.xlsx)
        self.salvar_portal([pedido()])

    def salvar_portal(self, ps):
        self.json.write_text(json.dumps({'PEDIDOS': ps, 'outro': 42}), encoding='utf-8')

    def portal(self):
        return json.loads(self.json.read_text(encoding='utf-8'))['PEDIDOS']

    def excel(self, celula, valor):
        w = openpyxl.load_workbook(self.xlsx)
        w['PEDIDOS'][celula] = valor
        w.save(self.xlsx)
        w.close()

    def rodar(self, gravar=True, **kw):
        return SP.sincronizar(str(self.xlsx), str(self.json), str(self.dados),
                              gravar=gravar, dormir=lambda _: None, **kw)

    def valor(self, celula):
        w = openpyxl.load_workbook(self.xlsx)
        try:
            return w['PEDIDOS'][celula].value
        finally:
            w.close()


@pytest.fixture
def c(tmp_path):
    return Cenario(tmp_path)


def test_primeira_rodada_e_previa_nao_alteram_dados(c):
    a, b = c.xlsx.read_bytes(), c.json.read_bytes()
    assert c.rodar(False)['primeira_vez']
    assert not list(c.dados.iterdir())
    assert c.rodar()['primeira_vez']
    assert c.xlsx.read_bytes() == a and c.json.read_bytes() == b
    assert c.rodar()['acoes'] == []
    assert c.xlsx.read_bytes() == a and c.json.read_bytes() == b


def test_novo_portal_acrescenta_linha_e_marca_visto(c):
    c.rodar()
    c.salvar_portal([pedido(), pedido(2)])
    assert len(c.rodar()['acoes']) == 1
    assert c.valor('C3') == 2
    estado = json.loads((c.dados / 'cadastro_pedidos.json').read_text(encoding='utf-8'))
    assert chave(pedido(2)) in estado['vistos']
    cad = cadastrar(str(c.xlsx), ArquivoPedidosConcorrente(str(c.json)), str(c.dados),
                    gravar=True, dormir=lambda _: None, progresso=lambda _: None)
    assert cad['gravados'] == 0 and len(c.portal()) == 2
    assert c.rodar()['acoes'] == []


def test_novo_excel_entra_no_portal(c):
    c.rodar()
    w = openpyxl.load_workbook(c.xlsx)
    w['PEDIDOS'].append(list(pedido(2).values()))
    w.save(c.xlsx)
    c.rodar()
    assert len(c.portal()) == 2
    assert 'NF (MAESTRO)' not in c.portal()[-1]


def test_excluido_excel_entre_cadastro_e_rodada_diaria_nao_volta(c):
    c.rodar()
    def cadastro():
        return cadastrar(str(c.xlsx), ArquivoPedidosConcorrente(str(c.json)), str(c.dados),
                         gravar=True, dormir=lambda _: None, progresso=lambda _: None)
    cadastro()
    w = openpyxl.load_workbook(c.xlsx)
    w['PEDIDOS'].append(list(pedido(2).values()))
    w.save(c.xlsx)
    assert cadastro()['gravados'] == 1
    w = openpyxl.load_workbook(c.xlsx)
    w['PEDIDOS'].delete_rows(3)
    w.save(c.xlsx)
    r = c.rodar()
    assert r['removidos'] and not r['acoes'] and c.valor('C3') is None


def test_registro_de_erro_nao_atropela_estado_compartilhado(c):
    from cruzar_nf.cadastro_pedidos import registrar_erro
    from cruzar_nf.sincronizar import trava, SincronizacaoEmAndamento
    with trava(c.dados, 'Outra rotina'):
        with pytest.raises(SincronizacaoEmAndamento):
            registrar_erro(str(c.dados), 'erro')
    assert not (c.dados / 'cadastro_pedidos.json').exists()


@pytest.mark.parametrize('campo,celula,valor', [('STATUS', 'I2', 'ENTREGUE'),
    ('DAV', 'K2', '123'), ('FATURAMENTO', 'H2', 'FATURADO'),
    ('VALOR ', 'D2', 123.45), ('DATA DE ENTREGA', 'F2', '12/11/2026')])
def test_mudanca_so_excel(c, campo, celula, valor):
    c.rodar()
    c.excel(celula, valor)
    c.rodar()
    assert c.portal()[0][campo] == valor


def test_mudanca_so_portal_preserva_xml_nao_autorizado(c):
    c.rodar()
    antes = c.xlsx.read_bytes()
    c.salvar_portal([pedido(**{'VALOR ': 222.50, 'DATA DE ENTREGA': '20/12/2026'})])
    c.rodar()
    assert c.valor('D2') == 222.5
    assert c.valor('F2').date() == dt.date(2026, 12, 20)
    with zipfile.ZipFile(io.BytesIO(antes)) as a, zipfile.ZipFile(c.xlsx) as b:
        parte = parte_da_aba(antes, 'PEDIDOS')
        assert all(a.read(n) == b.read(n) for n in a.namelist() if n != parte)
    assert c.rodar()['acoes'] == []


def test_conflito_bloqueia_item_inteiro(c):
    c.rodar()
    c.excel('I2', 'EXCEL')
    c.salvar_portal([pedido(STATUS='PORTAL', DAV='999')])
    r = c.rodar()
    assert r['conflitos'] and not r['acoes']
    assert c.valor('K2') is None and c.portal()[0]['STATUS'] == 'PORTAL'


def test_divergencia_inicial_nao_escolhe_vencedor(c):
    c.salvar_portal([pedido(STATUS='PORTAL')])
    c.rodar()
    c.excel('I2', 'OUTRO')
    assert c.rodar()['conflitos'] and c.portal()[0]['STATUS'] == 'PORTAL'


@pytest.mark.parametrize('lado', ['excel', 'portal'])
def test_exclusao_nao_ressuscita(c, lado):
    c.rodar()
    if lado == 'excel':
        w = openpyxl.load_workbook(c.xlsx)
        w['PEDIDOS'].delete_rows(2)
        w.save(c.xlsx)
    else:
        c.salvar_portal([])
    assert c.rodar()['removidos']
    assert c.rodar()['acoes'] == []
    assert (c.valor('C2') is None) if lado == 'excel' else c.portal() == []


def test_historico_so_portal_nao_e_importado(c):
    c.salvar_portal([pedido(), pedido(2)])
    c.rodar()
    assert c.rodar()['acoes'] == [] and c.valor('C3') is None


def test_ignorado_nao_e_sincronizado(c):
    c.rodar()
    c.salvar_portal([pedido(), pedido(2)])
    (c.dados / 'conferencia_ignorar.json').write_text(json.dumps(
        [{'chave': chave(pedido(2)), 'motivo': 'cancelado'}]), encoding='utf-8')
    assert c.rodar()['acoes'] == []


def test_ambiguo_nao_escreve(c):
    c.rodar()
    c.salvar_portal([pedido(STATUS='X'), pedido(STATUS='Y')])
    r = c.rodar()
    assert r['ambiguas'] and not r['acoes']


def test_correcao_identidade_so_relata(c):
    c.rodar()
    c.excel('C2', 99)
    r = c.rodar()
    assert r['conflitos'] and not r['acoes']


def test_correcao_manual_nas_duas_pontas_resolve_identidade(c):
    c.rodar()
    c.excel('C2', 99)
    c.salvar_portal([pedido(PEDIDO=99)])
    assert not c.rodar()['conflitos']
    c.excel('I2', 'ENTREGUE')
    c.rodar()
    assert c.portal()[0]['STATUS'] == 'ENTREGUE'


def test_aberta_adia_sem_mudar_foto(c, monkeypatch):
    c.rodar()
    foto = (c.dados / SP.ARQ_ESTADO).read_bytes()
    c.salvar_portal([pedido(STATUS='NOVO')])
    monkeypatch.setattr(SP.NP, 'esta_aberta', lambda _: 'Planilha aberta')
    r = c.rodar()
    assert r['adiado'] and c.valor('I2') == 'ABERTO'
    assert (c.dados / SP.ARQ_ESTADO).read_bytes() == foto


def test_excel_muda_antes_da_troca(c):
    c.rodar()
    c.salvar_portal([pedido(STATUS='NOVO')])
    with pytest.raises(GravacaoRecusada):
        c.rodar(antes_de_trocar=lambda: c.excel('B2', 'FOB'))
    assert c.valor('B2') == 'FOB' and c.valor('I2') == 'ABERTO'


def test_gerenciador_salva_por_cima_recompara(c):
    c.rodar()
    c.excel('I2', 'EXCEL')
    vezes = []
    def dormir(_):
        if not vezes:
            vezes.append(1)
            c.salvar_portal([pedido(), pedido(3)])
    SP.sincronizar(str(c.xlsx), str(c.json), str(c.dados), gravar=True, dormir=dormir)
    assert c.portal()[0]['STATUS'] == 'EXCEL'
    assert len(c.portal()) == 2


def test_tabela_filtro_estilo_e_outra_aba_preservados(tmp_path):
    c = Cenario(tmp_path, tabela=True)
    c.rodar()
    c.salvar_portal([pedido(), pedido(2)])
    c.rodar()
    w = openpyxl.load_workbook(c.xlsx)
    assert w['PEDIDOS'].tables['Pedidos'].ref == 'A1:N3'
    assert w['PEDIDOS'].auto_filter.ref == 'A1:N3'
    assert w['PEDIDOS']['C3'].font.bold
    assert w['OUTRA']['A1'].value == '=1+2'
    w.close()


def test_formula_destino_nao_e_substituida(c):
    c.excel('D2', '=50+50')
    c.rodar()
    c.salvar_portal([pedido(**{'VALOR ': 999})])
    assert c.rodar()['pendentes']
    assert c.valor('D2') == '=50+50'


def test_correspondencia_sem_produto_nao_autoriza_escrita():
    from cruzar_nf.cadastro_pedidos import correspondentes
    assert correspondentes({'PEDIDO': 1}, [pedido()]) == []
    assert correspondentes(pedido(), [{'PEDIDO': 1}]) == []


def test_retoma_apos_queda_entre_excel_e_foto(c, monkeypatch):
    c.rodar()
    c.salvar_portal([pedido(), pedido(2)])
    original = SP._gravar_portal
    monkeypatch.setattr(SP, '_gravar_portal', lambda *a: (_ for _ in ()).throw(RuntimeError('queda')))
    with pytest.raises(RuntimeError, match='queda'):
        c.rodar()
    assert c.valor('C3') == 2
    monkeypatch.setattr(SP, '_gravar_portal', original)
    assert c.rodar()['acoes'] == []
    assert c.valor('C4') is None
    assert 'pendente' not in json.loads((c.dados / SP.ARQ_ESTADO).read_text(encoding='utf-8'))


def test_retoma_novo_excel_apos_marcacao_visto(c, monkeypatch):
    c.rodar()
    w = openpyxl.load_workbook(c.xlsx)
    w['PEDIDOS'].append(list(pedido(2).values()))
    w.save(c.xlsx)
    original = SP._gravar_portal
    monkeypatch.setattr(SP, '_gravar_portal', lambda *a: (_ for _ in ()).throw(RuntimeError('queda')))
    with pytest.raises(RuntimeError):
        c.rodar()
    monkeypatch.setattr(SP, '_gravar_portal', original)
    c.rodar()
    assert len(c.portal()) == 2


def test_edicao_legitima_durante_verificacao_nao_e_sobrescrita(c):
    c.rodar()
    c.excel('I2', 'EXCEL')
    with pytest.raises(SP.ConflitoGravacao):
        SP.sincronizar(str(c.xlsx), str(c.json), str(c.dados), gravar=True,
                       dormir=lambda _: c.salvar_portal([pedido(STATUS='HUMANO')]))
    assert c.portal()[0]['STATUS'] == 'HUMANO'
    assert c.rodar()['conflitos']


def test_mesma_mudanca_nos_dois_lados_avanca_foto(c):
    c.rodar()
    c.excel('I2', 'ENTREGUE')
    c.salvar_portal([pedido(STATUS='ENTREGUE')])
    assert c.rodar()['acoes'] == []
    c.excel('I2', 'FINAL')
    c.rodar()
    assert c.portal()[0]['STATUS'] == 'FINAL'


def test_diario_respeita_janela_e_repeticao(c):
    assert c.rodar(agora=dt.datetime(2026, 10, 9, 10), diario=True)['sem_mudanca']
    assert not list(c.dados.iterdir())
    c.rodar(agora=dt.datetime(2026, 10, 9, 12, 30), diario=True)
    assert c.rodar(agora=dt.datetime(2026, 10, 9, 12, 40), diario=True)['sem_mudanca']


def test_estado_do_portal_inclui_novo_relatorio(c):
    from cruzar_nf.maestro import registrar
    from tests.test_sincronizar import Sio, Log
    c.rodar()
    sio = Sio()
    h = registrar(sio, {'caminho_banco_dados': str(c.dados)}, Log(), manager=object(), em_thread=False)
    h['estado']()
    assert sio.ultimo('retorno_sync_nf_estado')['estado']['sync_planilha_portal']['primeira_vez']


def test_dois_novos_portal_com_mesma_chave_nao_duplicam_excel(c):
    c.rodar()
    c.salvar_portal([pedido(), pedido(2), pedido(2)])
    r = c.rodar()
    assert r['ambiguas'] and not r['acoes'] and c.valor('C3') is None


def test_datas_legadas_americanas_sao_comparadas_com_excel(c):
    c.excel('F2', dt.datetime(2025, 7, 3))
    c.salvar_portal([pedido(**{'DATA DE ENTREGA': '7/3/2025'})])
    assert not c.rodar()['conferencia']['campos_divergentes']
    assert not c.rodar()['conflitos']
    c.salvar_portal([pedido(**{'DATA DE ENTREGA': '8/4/2025'})])
    c.rodar()
    assert c.valor('F2').date() == dt.date(2025, 8, 4)
    c.excel('F2', dt.datetime(2025, 9, 5))
    c.rodar()
    assert c.portal()[0]['DATA DE ENTREGA'] == '05/09/2025'
    assert c.rodar()['acoes'] == []


def test_email_editado_nao_mantem_link_para_destinatario_antigo(c):
    w = openpyxl.load_workbook(c.xlsx)
    w['PEDIDOS']['M2'].hyperlink = 'mailto:ana@vale.com'
    w['PEDIDOS']['A2'].hyperlink = 'https://example.com'
    w.save(c.xlsx)
    c.rodar()
    c.salvar_portal([pedido(**{'EMAIL REQUSITAN': 'bia@vale.com'})])
    c.rodar()
    w = openpyxl.load_workbook(c.xlsx)
    assert w['PEDIDOS']['M2'].value == 'bia@vale.com'
    assert w['PEDIDOS']['M2'].hyperlink is None
    assert w['PEDIDOS']['A2'].hyperlink.target == 'https://example.com'
    w.close()


def test_mesclagem_nao_e_sobrescrita(c):
    w = openpyxl.load_workbook(c.xlsx)
    w['PEDIDOS'].merge_cells('M2:M3')
    w.save(c.xlsx)
    c.rodar()
    c.salvar_portal([pedido(**{'EMAIL REQUSITAN': 'bia@vale.com'})])
    assert c.rodar()['pendentes']
    assert c.valor('M2') == 'ana@vale.com'


def test_validacao_e_formatacao_condicional_acompanham_linha_nova(c):
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.formatting.rule import CellIsRule
    w = openpyxl.load_workbook(c.xlsx)
    s = w['PEDIDOS']
    dv = DataValidation(type='list', formula1='"ABERTO,ENTREGUE"')
    s.add_data_validation(dv)
    dv.add('I2:I2')
    s.conditional_formatting.add('D2:D2', CellIsRule(operator='greaterThan', formula=['100']))
    w.save(c.xlsx)
    c.rodar()
    c.salvar_portal([pedido(), pedido(2)])
    c.rodar()
    w = openpyxl.load_workbook(c.xlsx)
    s = w['PEDIDOS']
    assert str(s.data_validations.dataValidation[0].sqref) == 'I2:I3'
    assert any(str(cf.sqref) == 'D2:D3' for cf in s.conditional_formatting)
    w.close()
