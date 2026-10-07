"""NF do portal gravada na planilha dos estagiários (cruzar_nf/nf_planilha.py).

O que mais importa: não estragar a planilha. Só a coluna N da aba PEDIDOS muda;
qualquer outra parte do arquivo continua byte a byte igual; planilha aberta no
Excel, ou que mudou no meio, não é gravada.
"""

from __future__ import annotations

import io
import json
import sys
import zipfile
from pathlib import Path

import openpyxl
import pytest
from openpyxl.styles import Font

from cruzar_nf import nf_planilha as NP
from cruzar_nf.nf_planilha import TITULO, editar_aba, esta_aberta, gravar_nf_na_planilha

CAB = ["CIDADE", "FRETE", "PEDIDO", "VALOR ", "PRODUTO", "DATA DE ENTREGA", "NMR DA RFQ",
       "FATURAMENTO", "STATUS", "Nº NOTA FISCAL", "DAV", "REQUISITANTE ", "EMAIL REQUSITAN"]


def criar_planilha(arq, linhas):
    """Parecida com a real: estilos, hyperlink de e-mail, coluna oculta, mesclada,
    filtro, painel congelado e uma segunda aba."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PEDIDOS"
    ws.append(CAB)
    for l in linhas:
        ws.append(l)
    for r in range(2, ws.max_row + 1):
        ws.cell(r, 10).font = Font(bold=True, color="FF0000")
        ws.cell(r, 13).hyperlink = f"mailto:{ws.cell(r, 13).value}"
    ws.column_dimensions["H"].hidden = True
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "L1:L3"
    if ws.max_row >= 3:
        ws.merge_cells("M2:M3")
    ws["N2"].font = Font(italic=True)                # célula formatada e vazia, como na real
    cot = wb.create_sheet("COTAÇÃO")
    for i in range(300):
        cot.append([f"COT{i}", "01/01/2026", "ITEM", i])
    wb.save(arq)


def lin(pedido, produto, nf_planilha=None, valor=100.0):
    return ["ITABIRA - MG", "EXW", pedido, valor, produto, "10/11/2026", None, None, None,
            nf_planilha, None, "LEANDRO", f"x{pedido}@vale.com"]


def ped(pedido, produto, nf=None, valor=100.0):
    d = {"CIDADE": "ITABIRA - MG", "PEDIDO": pedido, "VALOR ": valor, "PRODUTO": produto}
    if nf:
        d["Nº NOTA FISCAL"] = nf
    return d


class Cenario:
    def __init__(self, tmp_path, linhas, pedidos):
        self.pasta = tmp_path / "Publico"
        self.pasta.mkdir()
        self.xlsx = self.pasta / "PLANILHA.xlsx"
        criar_planilha(self.xlsx, linhas)
        self.pedidos = tmp_path / "PEDIDOS.json"
        self.pedidos.write_text(json.dumps({"PEDIDOS": pedidos}, ensure_ascii=False), encoding="utf-8")
        self.dados = tmp_path / "dados"
        self.dados.mkdir()
        self.original = self.xlsx.read_bytes()

    def rodar(self, gravar=True, **kw):
        return gravar_nf_na_planilha(str(self.xlsx), str(self.pedidos), str(self.dados), gravar=gravar,
                                     progresso=lambda _m: None, **kw)

    def valores(self, col):
        wb = openpyxl.load_workbook(self.xlsx, read_only=True)
        try:
            return [r[0] for r in wb["PEDIDOS"].iter_rows(min_col=col, max_col=col, values_only=True)]
        finally:
            wb.close()


@pytest.fixture
def cen(tmp_path):
    return Cenario(tmp_path, [lin(1, "MONITOR"), lin(2, "TECLADO"), lin(3, "MOUSE")],
                   [ped(1, "MONITOR 24", nf="15730"), ped(2, "TECLADO", nf="15020 / 15031"),
                    ped(3, "MOUSE")])


def partes(dados):
    with zipfile.ZipFile(io.BytesIO(dados)) as z:
        return {i.filename: z.read(i.filename) for i in z.infolist()}


# -- gravação ------------------------------------------------------------------------

def test_grava_so_a_coluna_n_e_o_resto_fica_byte_a_byte_igual(cen):
    res = cen.rodar()
    assert res["gravadas"] == 2 and res["gravou"] is True
    assert cen.valores(14) == [TITULO, 15730, "15020 / 15031", None]
    antes, depois = partes(cen.original), partes(cen.xlsx.read_bytes())
    assert list(antes) == list(depois)                                  # mesmas partes, mesma ordem
    aba = NP.parte_da_aba(cen.original, "PEDIDOS")
    assert all(antes[k] == depois[k] for k in antes if k != aba)
    wb0 = openpyxl.load_workbook(io.BytesIO(cen.original), read_only=True)
    for col in range(1, 14):                                            # nenhuma outra coluna mudou
        assert [r[0] for r in wb0["PEDIDOS"].iter_rows(min_col=col, max_col=col, values_only=True)] \
            == cen.valores(col)
    wb0.close()
    xml0, xml1 = antes[aba].decode(), depois[aba].decode()
    for bloco in ("<hyperlinks>", "<mergeCells", "<autoFilter", "<cols>", "<pane "):
        trecho = xml0[xml0.index(bloco):xml0.index(">", xml0.index(bloco)) + 1]
        assert trecho in xml1


def test_mantem_hyperlinks_formatos_e_abre_no_openpyxl_completo(cen):
    cen.rodar()
    wb = openpyxl.load_workbook(cen.xlsx)
    ws = wb["PEDIDOS"]
    assert ws["M2"].hyperlink.target == "mailto:x1@vale.com"
    assert ws["J2"].font.b and ws.column_dimensions["H"].hidden
    assert ws["N2"].font.i                           # a célula formatada que já existia manteve o estilo
    assert ws["N3"].font.b                           # a nova copia o estilo da coluna de NF (J)
    assert "M2:M3" in [str(m) for m in ws.merged_cells.ranges] and wb["COTAÇÃO"].max_row == 300


def test_previa_nao_toca_no_arquivo(cen):
    res = cen.rodar(gravar=False)
    assert res["a_gravar"] == 2 and res["gravadas"] == 0
    assert cen.xlsx.read_bytes() == cen.original
    assert not list(cen.dados.glob("**/*.xlsx"))


def test_segunda_vez_nao_regrava_nada(cen):
    cen.rodar()
    depois = cen.xlsx.read_bytes()
    res = cen.rodar()
    assert res["a_gravar"] == 0 and res["gravou"] is False
    assert cen.xlsx.read_bytes() == depois


def test_nf_que_mudou_no_portal_e_atualizada(cen):
    cen.rodar()
    cen.pedidos.write_text(json.dumps({"PEDIDOS": [ped(1, "MONITOR", nf="15999")]}), encoding="utf-8")
    res = cen.rodar()
    assert res["atualizadas"] == 1 and cen.valores(14)[1] == 15999


def test_linha_com_dois_pedidos_de_nf_diferente_nao_e_gravada(tmp_path):
    c = Cenario(tmp_path, [lin(1, "MONITOR")], [ped(1, "MONITOR", nf="10"), ped(1, "MONITOR 24", nf="11")])
    res = c.rodar()
    assert res["gravadas"] == 0 and res["ambiguas"][0]["linha"] == 2


def test_backup_antes_de_gravar(cen):
    res = cen.rodar()
    backups = list((cen.dados / "backups_planilha").glob("*.xlsx"))
    assert len(backups) == 1 and backups[0].read_bytes() == cen.original
    assert res["backup"] == str(backups[0])


def test_coluna_n_usada_para_outra_coisa_nao_e_tocada(tmp_path):
    c = Cenario(tmp_path, [lin(1, "MONITOR")], [ped(1, "MONITOR", nf="10")])
    wb = openpyxl.load_workbook(c.xlsx)
    wb["PEDIDOS"]["N1"] = "OBSERVAÇÕES"
    wb.save(c.xlsx)
    antes = c.xlsx.read_bytes()
    with pytest.raises(NP.GravacaoRecusada, match="OBSERVAÇÕES"):
        c.rodar()
    assert c.xlsx.read_bytes() == antes


def test_cadastro_automatico_nao_leva_a_coluna_nf_maestro_para_o_portal():
    from cruzar_nf.cadastro_pedidos import _sem_colunas_do_robo
    assert _sem_colunas_do_robo({"PEDIDO": 1, TITULO: 15730}) == {"PEDIDO": 1}


# -- segurança: aberta, mudou no meio, falha na troca -------------------------------------

def segurar_como_o_excel(caminho):
    """Abre como o Excel abre: lê e grava, deixando os outros só lerem."""
    import ctypes
    from ctypes import wintypes
    k = ctypes.WinDLL("kernel32", use_last_error=True)
    k.CreateFileW.restype = wintypes.HANDLE
    k.CreateFileW.argtypes = [wintypes.LPCWSTR, wintypes.DWORD, wintypes.DWORD, wintypes.LPVOID,
                              wintypes.DWORD, wintypes.DWORD, wintypes.HANDLE]
    h = k.CreateFileW(str(caminho), 0x80000000 | 0x40000000, 0x1, None, 3, 0x80, None)
    assert h not in (None, wintypes.HANDLE(-1).value)
    return lambda: k.CloseHandle(wintypes.HANDLE(h))


windows = pytest.mark.skipif(sys.platform != "win32", reason="trava de arquivo do Windows")


@windows
def test_planilha_aberta_no_excel_nao_e_gravada_e_diz_quem_abriu(cen):
    (cen.pasta / "~$PLANILHA.xlsx").write_bytes(bytes([6]) + b"GUSTAV" + b"\x00" * 40)
    soltar = segurar_como_o_excel(cen.xlsx)
    try:
        assert "GUSTAV" in esta_aberta(str(cen.xlsx))
        with pytest.raises(NP.GravacaoRecusada, match="aberta"):
            cen.rodar()
    finally:
        soltar()
    assert cen.xlsx.read_bytes() == cen.original
    assert esta_aberta(str(cen.xlsx)) is None        # trava velha do Excel sozinha não impede


@windows
def test_alguem_abre_no_ultimo_instante(cen):
    soltar = []
    with pytest.raises(NP.GravacaoRecusada, match="aberta"):
        cen.rodar(_antes_de_trocar=lambda: soltar.append(segurar_como_o_excel(cen.xlsx)))
    soltar[0]()
    assert cen.xlsx.read_bytes() == cen.original
    assert [p.name for p in cen.pasta.iterdir()] == ["PLANILHA.xlsx"]   # sem temporário largado


def test_planilha_salva_no_meio_nao_e_sobrescrita(cen):
    def estagiario_salva():
        wb = openpyxl.load_workbook(cen.xlsx)
        wb["PEDIDOS"].append(lin(4, "NOVO"))
        wb.save(cen.xlsx)
    with pytest.raises(NP.GravacaoRecusada, match="mudou"):
        cen.rodar(_antes_de_trocar=estagiario_salva)
    assert cen.valores(3)[-1] == 4                    # o que o estagiário salvou continua lá


def test_falha_na_troca_deixa_o_original_e_nao_larga_temporario(cen, monkeypatch):
    def troca_falha(_a, _b):
        raise PermissionError("acesso negado")
    monkeypatch.setattr(NP.os, "replace", troca_falha)
    with pytest.raises(NP.GravacaoRecusada, match="aberta"):
        cen.rodar()
    assert cen.xlsx.read_bytes() == cen.original
    assert [p.name for p in cen.pasta.iterdir()] == ["PLANILHA.xlsx"]


def test_conferencia_depois_de_gravar_pega_arquivo_diferente(cen, monkeypatch):
    real = NP._trocar

    def troca_e_alguem_grava_por_cima(tmp, destino):
        real(tmp, destino)
        Path(destino).write_bytes(cen.original)
    monkeypatch.setattr(NP, "_trocar", troca_e_alguem_grava_por_cima)
    with pytest.raises(NP.GravacaoRecusada, match="confer"):
        cen.rodar()


# -- edição do XML (casos que o openpyxl não gera) ----------------------------------------

def test_edita_linha_autofechada_celula_formatada_e_ordem_das_colunas():
    xml = ('<worksheet><dimension ref="A1:M3"/><sheetData>'
           '<row r="1" spans="1:13"><c r="A1" t="s"><v>0</v></c><c r="J1" s="7"/></row>'
           '<row r="2" spans="1:13"><c r="J2" s="8"><v>1</v></c><c r="N2" s="9"/><c r="P2" s="3"/></row>'
           '<row r="3" spans="1:13" ht="15"/>'
           '</sheetData></worksheet>')
    novo = editar_aba(xml, {1: TITULO, 2: "15020 / 15031", 3: "15730"})
    assert '<c r="J1" s="7"/><c r="N1" s="7" t="inlineStr"><is><t>NF (MAESTRO)</t></is></c></row>' in novo
    assert '<c r="N2" s="9" t="inlineStr"><is><t>15020 / 15031</t></is></c><c r="P2" s="3"/>' in novo
    assert '<row r="3" spans="1:14" ht="15"><c r="N3"><v>15730</v></c></row>' in novo
    assert '<dimension ref="A1:N3"/>' in novo


def test_texto_com_caracteres_especiais_e_escapado():
    xml = '<worksheet><sheetData><row r="2"><c r="A2"/></row></sheetData></worksheet>'
    assert "<t>A &amp; B &lt;1&gt;</t>" in editar_aba(xml, {2: "A & B <1>"})


def test_linha_que_nao_existe_no_xml_e_erro():
    with pytest.raises(NP.GravacaoRecusada):
        editar_aba('<worksheet><sheetData><row r="2"/></sheetData></worksheet>', {5: "1"})


# -- a conferência pega edição errada ------------------------------------------------------

def test_conferencia_pega_mudanca_em_outra_coluna(cen):
    aba = NP.parte_da_aba(cen.original, "PEDIDOS")
    with zipfile.ZipFile(io.BytesIO(cen.original)) as z:
        xml = z.read(aba).decode()
    estragado = NP._reempacotar(cen.original, aba, NP.editar_aba(xml, {2: "999"}, coluna="C").encode())
    with pytest.raises(NP.GravacaoRecusada, match="linha 2"):
        NP.conferir(cen.original, estragado, {})


def test_edicao_errada_e_barrada_antes_de_gravar_e_na_previa(cen, monkeypatch):
    real = NP.editar_aba
    monkeypatch.setattr(NP, "editar_aba", lambda xml, valores: real(real(xml, valores), {2: "X"}, coluna="A"))
    for gravar in (False, True):
        with pytest.raises(NP.GravacaoRecusada, match="conferência"):
            cen.rodar(gravar=gravar)
    assert cen.xlsx.read_bytes() == cen.original
    assert not (cen.dados / "backups_planilha").exists()      # barrou antes até do backup


def test_mesma_rfq_e_produto_em_pedidos_diferentes_vale_o_de_mesmo_numero(tmp_path):
    # na planilha real: várias linhas da mesma RFQ, mesmo produto, uma NF por pedido
    c = Cenario(tmp_path, [lin(1, "MONITOR"), lin(2, "MONITOR")],
                [ped(1, "MONITOR", nf="10"), ped(2, "MONITOR", nf="11")])
    res = c.rodar()
    assert res["ambiguas"] == [] and c.valores(14)[1:] == [10, 11]


# -- botão do portal (via serviço de NF) ---------------------------------------------------

def registrar_no_portal(cen):
    from tests.test_sincronizar import Log, Sio
    from cruzar_nf.maestro import registrar
    sio = Sio()
    h = registrar(sio, {"caminho_banco_dados": str(cen.dados), "caminho_pedidos": str(cen.pedidos),
                        "caminho_planilha": str(cen.xlsx)}, Log(), manager=object(), em_thread=False)
    return sio, h


def test_botao_previa_e_gravar_pelo_portal(cen):
    sio, h = registrar_no_portal(cen)
    h["nf_planilha"]({"clientId": "x", "gravar": False})
    r = sio.ultimo("retorno_nf_planilha")
    assert r["sucesso"] and r["clientId"] == "x" and r["resultado"]["a_gravar"] == 2
    assert cen.xlsx.read_bytes() == cen.original
    assert any(e == "progresso_nf_planilha" for e, _ in sio.emitidos)
    h["nf_planilha"]({"clientId": "x", "gravar": True})
    assert sio.ultimo("retorno_nf_planilha")["resultado"]["gravadas"] == 2
    assert cen.valores(14)[1] == 15730


@windows
def test_botao_com_planilha_aberta_responde_erro_claro(cen):
    sio, h = registrar_no_portal(cen)
    soltar = segurar_como_o_excel(cen.xlsx)
    try:
        h["nf_planilha"]({"clientId": "x", "gravar": True})
    finally:
        soltar()
    r = sio.ultimo("retorno_nf_planilha")
    assert r["sucesso"] is False and "aberta" in r["erro"]
    assert cen.xlsx.read_bytes() == cen.original


def test_estado_leva_o_resumo_da_conferencia(cen):
    from cruzar_nf.planilha_portal import gerar
    sio, h = registrar_no_portal(cen)
    h["estado"]({"clientId": "y"})
    assert sio.ultimo("retorno_sync_nf_estado")["estado"]["conferencia_planilha"] is None
    gerar(str(cen.xlsx), str(cen.pedidos), str(cen.dados))
    h["estado"]({"clientId": "y"})
    conf = sio.ultimo("retorno_sync_nf_estado")["estado"]["conferencia_planilha"]
    assert conf["planilha_sem_portal"] == 0 and conf["nf_so_no_portal"] == 2
