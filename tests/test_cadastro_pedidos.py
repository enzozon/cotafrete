"""Cadastro automático: planilha Excel dos estagiários -> PEDIDOS.json do Maestro
(cruzar_nf/cadastro_pedidos.py)."""

from __future__ import annotations

import datetime as dt
import json
import os
from pathlib import Path

import openpyxl
import pytest

from cruzar_nf.arquivo_pedidos import ArquivoPedidosConcorrente
from cruzar_nf.cadastro_pedidos import cadastrar, chave, ler_planilha

CABECALHO = ["CIDADE", "FRETE", "PEDIDO", "VALOR ", "PRODUTO", "DATA DE ENTREGA", "NMR DA RFQ",
             "FATURAMENTO", "STATUS", "Nº NOTA FISCAL", "DAV", "REQUISITANTE ", "EMAIL REQUSITAN"]


def linha(pedido, produto="13182678 || FORNO MICROONDAS 31L", cidade="Barão de Cocais, MG",
          valor=650.4, rfq=284014, req="LEANDRO", entrega=dt.datetime(2026, 11, 16)):
    return [cidade, "EXW", pedido, valor, produto, entrega, rfq, None, None, None, None, req,
            "fulano@vale.com"]


def escrever_planilha(arq, linhas):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PEDIDOS"
    ws.append(CABECALHO)
    for l in linhas:
        ws.append(l)
    wb.create_sheet("COTAÇÃO").append(["COTAÇÃO"])
    wb.save(arq)
    # a data de modificação anda a cada salvamento, como no Excel
    escrever_planilha.n = getattr(escrever_planilha, "n", 0) + 10
    t = dt.datetime.now().timestamp() + escrever_planilha.n
    os.utime(arq, (t, t))


def escrever_pedidos(arq, pedidos):
    with open(arq, "w", encoding="utf-8") as f:
        json.dump({"PEDIDOS": pedidos}, f, indent=1, ensure_ascii=False)


def ler_pedidos(arq):
    return json.loads(Path(arq).read_text(encoding="utf-8"))["PEDIDOS"]


class Cenario:
    def __init__(self, tmp_path):
        self.xlsx = tmp_path / "PLANILHA.xlsx"
        self.pedidos = tmp_path / "PEDIDOS.json"
        self.dados = tmp_path / "dados"
        self.dados.mkdir()
        self.antigas = [linha(4513619890, "35002018 || REFRIGERADORES", rfq=278377),
                        linha(4513631320, "13509761 || CONDICIONADOR AR", rfq=281352)]
        escrever_planilha(self.xlsx, self.antigas)
        escrever_pedidos(self.pedidos, [{"PEDIDO": 4513619890, "PRODUTO": "35002018 || REFRIGERADORES",
                                         "NMR DA RFQ": 278377}])

    def rodar(self, gravar=True, dormir=lambda _s: None):
        manager = ArquivoPedidosConcorrente(str(self.pedidos), dormir=lambda _s: None)
        return cadastrar(str(self.xlsx), manager, str(self.dados), gravar=gravar,
                         progresso=lambda _m: None, dormir=dormir)


@pytest.fixture
def cenario(tmp_path):
    return Cenario(tmp_path)


def test_le_a_aba_pedidos_no_formato_do_portal(tmp_path):
    arq = tmp_path / "p.xlsx"
    escrever_planilha(arq, [linha(4513646471), [None] * 13, [None, "EXW"]])
    assert ler_planilha(str(arq)) == [{
        "CIDADE": "Barão de Cocais, MG", "FRETE": "EXW", "PEDIDO": 4513646471, "VALOR ": 650.4,
        "PRODUTO": "13182678 || FORNO MICROONDAS 31L", "DATA DE ENTREGA": "16/11/2026",
        "NMR DA RFQ": 284014, "REQUISITANTE ": "LEANDRO", "EMAIL REQUSITAN": "fulano@vale.com",
    }, {"FRETE": "EXW"}]


def test_chave_ignora_acento_espaco_e_tipo():
    a = {"PEDIDO": 4513646471, "NMR DA RFQ": "284014", "PRODUTO": "Forno  micro-ondas"}
    b = {"PEDIDO": "4513646471 ", "NMR DA RFQ": 284014, "PRODUTO": "FORNO MICROONDAS"}
    assert chave(a) == chave(b)


def test_primeira_vez_so_marca_o_que_ja_existe(cenario):
    antes = cenario.pedidos.read_text(encoding="utf-8")
    res = cenario.rodar()
    assert res["primeira_vez"] is True and res["gravados"] == 0
    assert cenario.pedidos.read_text(encoding="utf-8") == antes
    escrever_planilha(cenario.xlsx, cenario.antigas)    # salvou de novo sem mudar nada
    assert cenario.rodar()["novas"] == 0


def test_incompleta_antiga_nao_vira_aviso_eterno(cenario):
    sem_numero = linha(None, "BANDEJA PARA ACOMODAÇÃO", rfq=None)
    escrever_planilha(cenario.xlsx, cenario.antigas + [sem_numero])
    cenario.rodar()
    escrever_planilha(cenario.xlsx, cenario.antigas + [sem_numero, linha(4513646471)])
    res = cenario.rodar()
    assert res["incompletas"] == 0 and res["gravados"] == 1


def test_previa_nao_grava_nem_guarda_estado(cenario):
    cenario.rodar()
    escrever_planilha(cenario.xlsx, cenario.antigas + [linha(4513646471)])
    res = cenario.rodar(gravar=False)
    assert res["a_gravar"] == 1 and res["gravados"] == 0
    assert len(ler_pedidos(cenario.pedidos)) == 1
    assert cenario.rodar()["gravados"] == 1


def test_linha_nova_entra_no_portal_uma_vez_so(cenario):
    cenario.rodar()
    escrever_planilha(cenario.xlsx, cenario.antigas + [linha(4513646471)])
    assert cenario.rodar()["gravados"] == 1
    novo = ler_pedidos(cenario.pedidos)[-1]
    assert novo["PEDIDO"] == 4513646471 and novo["DATA DE ENTREGA"] == "16/11/2026"
    assert novo["VALOR "] == 650.4 and novo["REQUISITANTE "] == "LEANDRO"
    assert list((cenario.dados / "backups").glob("PEDIDOS antes do cadastro *.json"))
    assert cenario.rodar()["sem_mudanca"] is True
    escrever_planilha(cenario.xlsx, cenario.antigas + [linha(4513646471)])   # salvou de novo
    assert cenario.rodar()["gravados"] == 0
    assert len(ler_pedidos(cenario.pedidos)) == 2


def test_pedido_com_varios_itens_entra_inteiro(cenario):
    cenario.rodar()
    escrever_planilha(cenario.xlsx, cenario.antigas + [linha(4513700000, "ITEM A MONITOR"),
                                                       linha(4513700000, "ITEM B TECLADO")])
    assert cenario.rodar()["gravados"] == 2


def test_item_novo_de_pedido_que_ja_esta_no_portal_entra(cenario):
    cenario.rodar()
    escrever_planilha(cenario.xlsx, cenario.antigas + [linha(4513619890, "99999999 || OUTRO ITEM",
                                                             rfq=278377)])
    assert cenario.rodar()["gravados"] == 1
    assert len(ler_pedidos(cenario.pedidos)) == 2


def test_pedido_ja_cadastrado_a_mao_no_portal_nao_duplica(cenario):
    cenario.rodar()
    p = ler_pedidos(cenario.pedidos)
    escrever_pedidos(cenario.pedidos, p + [{"PEDIDO": "4513646471", "NMR DA RFQ": "284014",
                                            "PRODUTO": "13182678 || FORNO MICROONDAS 31L S/INF"}])
    escrever_planilha(cenario.xlsx, cenario.antigas + [linha(4513646471)])
    res = cenario.rodar()
    assert res["gravados"] == 0 and res["ja_no_portal"] == 1
    assert len(ler_pedidos(cenario.pedidos)) == 2


def test_correcao_de_digitacao_em_linha_antiga_nao_duplica(cenario):
    cenario.rodar()
    corrigida = [linha(4513619899, "35002018 || REFRIGERADORES", rfq=278377), cenario.antigas[1]]
    escrever_planilha(cenario.xlsx, corrigida)          # pedido mudou; RFQ e produto iguais
    assert cenario.rodar()["gravados"] == 0
    assert len(ler_pedidos(cenario.pedidos)) == 1


def test_linha_incompleta_espera_ficar_pronta(cenario):
    cenario.rodar()
    escrever_planilha(cenario.xlsx, cenario.antigas + [linha(4513646471, valor=None, req=None)])
    res = cenario.rodar()
    assert res["gravados"] == 0 and res["incompletas"] == 1
    escrever_planilha(cenario.xlsx, cenario.antigas + [linha(4513646471)])
    assert cenario.rodar()["gravados"] == 1


def test_apagado_no_portal_nao_volta(cenario):
    cenario.rodar()
    escrever_planilha(cenario.xlsx, cenario.antigas + [linha(4513646471)])
    cenario.rodar()
    escrever_pedidos(cenario.pedidos, ler_pedidos(cenario.pedidos)[:1])   # alguém apagou no portal
    escrever_planilha(cenario.xlsx, cenario.antigas + [linha(4513646471)])
    assert cenario.rodar()["gravados"] == 0
    assert len(ler_pedidos(cenario.pedidos)) == 1


def test_gerenciador_gravando_por_cima_faz_gravar_de_novo(cenario):
    cenario.rodar()
    escrever_planilha(cenario.xlsx, cenario.antigas + [linha(4513646471)])
    original = ler_pedidos(cenario.pedidos)
    vezes = []

    def gerenciador_salva_memoria_velha(_s):
        if not vezes:           # na espera da conferência, o gerenciador salva a memória velha
            vezes.append(1)
            escrever_pedidos(cenario.pedidos, original)
            t = dt.datetime.now().timestamp() + 600
            os.utime(cenario.pedidos, (t, t))

    res = cenario.rodar(dormir=gerenciador_salva_memoria_velha)
    assert vezes and res["gravados"] == 1
    assert [p["PEDIDO"] for p in ler_pedidos(cenario.pedidos)] == [4513619890, 4513646471]


def test_planilha_no_meio_do_salvamento_nao_estraga_o_estado(cenario):
    cenario.rodar()
    estado = (cenario.dados / "cadastro_pedidos.json").read_text(encoding="utf-8")
    cenario.xlsx.write_bytes(b"meio arquivo")
    with pytest.raises(ValueError, match="salvando"):
        cenario.rodar()
    assert (cenario.dados / "cadastro_pedidos.json").read_text(encoding="utf-8") == estado
    escrever_planilha(cenario.xlsx, cenario.antigas + [linha(4513646471)])
    assert cenario.rodar()["gravados"] == 1


def test_planilha_sem_aba_pedidos_da_erro_claro(tmp_path):
    arq = tmp_path / "x.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "OUTRA"
    wb.save(arq)
    with pytest.raises(ValueError, match="PEDIDOS"):
        ler_planilha(str(arq))


def test_o_gerenciador_do_servidor_mostra_o_pedido_sem_reiniciar(cenario):
    from tests.test_arquivo_pedidos import carregar_planilha_manager_do_servidor

    (cenario.pedidos.parent / "COTAÇÕES.json").write_text('{"COTAÇÃO": []}', encoding="utf-8")
    pm = carregar_planilha_manager_do_servidor(cenario.pedidos.parent)
    pm.iniciar()
    cenario.rodar()
    escrever_planilha(cenario.xlsx, cenario.antigas + [linha(4513646471)])
    cenario.rodar()

    pm.iniciar()                                         # o que a tela "Pedidos" faz ao carregar
    assert [p["PEDIDO"] for p in pm.pedidos["PEDIDOS"]] == [4513619890, 4513646471]
    # e um cadastro manual pelo portal depois não apaga o pedido que veio da planilha
    pm.adicionar_linhas_pedidos("PEDIDO", [{"pedido": 1, "itens": []}])
    assert len(ler_pedidos(cenario.pedidos)) == 3


def test_numero_corrigido_em_linha_sem_rfq_nao_duplica(cenario):
    cenario.rodar()
    escrever_planilha(cenario.xlsx, cenario.antigas + [linha(1234, "BANDEJA", cidade="CATU-BA",
                                                             valor=2442.49, rfq=None)])
    assert cenario.rodar()["gravados"] == 1
    escrever_planilha(cenario.xlsx, cenario.antigas + [linha(1243, "BANDEJA", cidade="CATU-BA",
                                                             valor=2442.49, rfq=None)])
    assert cenario.rodar()["gravados"] == 0
    assert len(ler_pedidos(cenario.pedidos)) == 2


def test_valor_no_formato_do_portal_conta_como_igual():
    from cruzar_nf.cadastro_pedidos import _valor
    assert _valor("3.219,70") == _valor(3219.7) == _valor("3219,70") == 3219.7
    assert _valor("5212.14") == 5212.14 and _valor("") is None


def test_linha_repetida_na_planilha_gera_aviso(cenario):
    cenario.rodar()
    escrever_planilha(cenario.xlsx, cenario.antigas + [linha(4513646471), linha(4513646471)])
    res = cenario.rodar()
    assert res["gravados"] == 1 and len(res["repetidas"]) == 1


# -- painel no portal ------------------------------------------------------------

def test_painel_mostra_ultimo_cadastro_e_ultima_verificacao(cenario):
    from cruzar_nf.cadastro_pedidos import estado_para_tela
    assert estado_para_tela(str(cenario.dados))["ativo"] is False
    cenario.rodar()
    escrever_planilha(cenario.xlsx, cenario.antigas + [linha(4513646471)])
    cenario.rodar()
    cenario.rodar()                                      # rodada sem mudança também conta como verificação
    tela = estado_para_tela(str(cenario.dados))
    assert tela["ativo"] is True and tela["erro"] is None
    assert tela["verificado_em"] and tela["ultimo_cadastro"]["gravados"] == 1
    assert tela["ultimo_cadastro"]["pedidos"] == ["4513646471 - 13182678 || FORNO MICROONDAS 31L"]
    assert tela["cadastrados_hoje"] == 1
    assert [h["gravados"] for h in tela["historico"]] == [1]   # a 1ª rodada (só marca) não aparece


def test_painel_mostra_o_erro_ate_a_proxima_rodada_boa(cenario):
    from cruzar_nf.cadastro_pedidos import estado_para_tela, main
    cenario.rodar()
    cenario.xlsx.write_bytes(b"meio arquivo")
    args = ["--planilha", str(cenario.xlsx), "--dados", str(cenario.dados), "--pedidos", str(cenario.pedidos)]
    assert main(args + ["--gravar"]) == 1
    assert "salvando" in estado_para_tela(str(cenario.dados))["erro"]["mensagem"]
    escrever_planilha(cenario.xlsx, cenario.antigas)
    assert main(args + ["--gravar"]) == 0
    assert estado_para_tela(str(cenario.dados))["erro"] is None


def test_previa_pela_linha_de_comando_nao_registra_nada(cenario):
    from cruzar_nf.cadastro_pedidos import main
    args = ["--planilha", str(cenario.xlsx), "--dados", str(cenario.dados), "--pedidos", str(cenario.pedidos)]
    assert main(args) == 0
    assert not (cenario.dados / "cadastro_pedidos.json").exists()


def test_o_servico_de_nf_manda_o_estado_do_cadastro_para_o_portal(cenario):
    from tests.test_sincronizar import Log, Sio
    from cruzar_nf.maestro import registrar
    cenario.rodar()
    sio = Sio()
    h = registrar(sio, {"caminho_banco_dados": str(cenario.dados)}, Log(), manager=object(), em_thread=False)
    h["estado"]({"clientId": "y"})
    r = sio.ultimo("retorno_sync_nf_estado")
    assert r["sucesso"] and r["estado"]["cadastro_pedidos"]["ativo"] is True


def test_estado_do_cadastro_quebrado_nao_derruba_o_estado_da_nf(cenario):
    from tests.test_sincronizar import Log, Sio
    from cruzar_nf.maestro import registrar
    (cenario.dados / "cadastro_pedidos.json").write_text("{quebrado", encoding="utf-8")
    sio = Sio()
    registrar(sio, {"caminho_banco_dados": str(cenario.dados)}, Log(), manager=object(),
              em_thread=False)["estado"]({"clientId": "y"})
    r = sio.ultimo("retorno_sync_nf_estado")
    assert r["sucesso"] and "erro" in r["estado"]["cadastro_pedidos"]["erro"]["mensagem"].lower()
